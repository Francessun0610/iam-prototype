#!/usr/bin/env python3
"""Extract the canonical role/permission model from Tatiana's workbook.

The workbook — "New OMS - Roles & Permissions.xlsx" — is the source of truth
for IAM v4.1 role definitions and permission grants. It is not committed (it
lives outside the repo), so this script freezes what it says into
`tests/workbook_model.json`, which `tests/test_workbook_parity.py` compares
the running application against.

    python3 tools/extract_workbook.py [path-to-workbook]

Reading rules, straight from the brief:
  * `x` in a role column is a grant; a blank cell is not.
  * Notes and Descriptions are prose. They never become grants.
  * `Core Planning` defines the Core Planning permissions; `v2. Agent` is
    authoritative for the two Agent roles and restates the Core Planning
    codes it needs so `Planning Agent User` can grant them. A restated code
    keeps its original `Core Planning` definition — the restatement adds
    grants, never a second permission.
  * `Agents` is a supporting sheet: granular capability descriptions behind
    the `v2. Agent` matrix. It is deliberately not read for grants.
"""
import json
import os
import sys

import openpyxl

DEFAULT_WB = os.path.expanduser("~/Downloads/New OMS - Roles & Permissions.xlsx")
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "tests", "workbook_model.json")

GRANT_SHEETS = ("Core Planning", "v2. Agent")
REFERENCE_SHEETS = ("Agents",)


def norm(v):
    return "" if v is None else str(v).strip()


def extract(path=DEFAULT_WB):
    wb = openpyxl.load_workbook(path, data_only=True)
    perms, order, grants, roles_order, issues = {}, [], {}, [], []

    for sheet in GRANT_SHEETS:
        rows = [[norm(c) for c in row] for row in wb[sheet].iter_rows(values_only=True)]
        hdr_i = next((i for i, r in enumerate(rows)
                      if any(c.lower() in ("function", "functions", "permission", "function code")
                             for c in r)), None)
        if hdr_i is None:
            issues.append([sheet, "no header row found"])
            continue
        low = [c.lower() for c in rows[hdr_i]]

        def col(*names):
            return next((low.index(n) for n in names if n in low), None)

        c_app, c_res = col("application", "app"), col("resource", "resources")
        c_fn = col("function", "functions", "permission", "function code")
        c_desc, c_note = col("description", "descriptions"), col("note", "notes")
        known = {c_app, c_res, c_fn, c_desc, c_note}

        role_cols = {}
        for j, name in enumerate(rows[hdr_i]):
            if j in known or not name:
                continue
            role_cols[j] = name
            if name not in roles_order:
                roles_order.append(name)
                grants[name] = set()

        last_app = last_res = ""
        for r in rows[hdr_i + 1:]:
            if not any(r):
                continue
            cell = lambda c: norm(r[c]) if c is not None and c < len(r) else ""
            last_app = cell(c_app) or last_app
            last_res = cell(c_res) or last_res
            fn = cell(c_fn)
            if not fn:
                continue
            if fn in perms:
                # A code restated on `v2. Agent` so an Agent role can grant it.
                # Keep the Core Planning definition; take only the grants.
                if perms[fn]["sheet"] != sheet:
                    issues.append([sheet, "restated", fn])
                else:
                    issues.append([sheet, "duplicate function code", fn])
            else:
                perms[fn] = {"code": fn, "application": last_app, "resource": last_res,
                             "description": cell(c_desc), "sheet": sheet}
                order.append(fn)
            for j, role in role_cols.items():
                mark = (norm(r[j]) if j < len(r) else "").lower()
                if mark == "x":
                    grants[role].add(fn)
                elif mark:
                    issues.append([sheet, "non-x mark ignored", fn, role, mark])

    for sheet in REFERENCE_SHEETS:
        if sheet not in wb.sheetnames:
            issues.append([sheet, "reference sheet missing"])

    return {
        "_source": os.path.basename(path) + " (Tatiana). Regenerate with tools/extract_workbook.py.",
        "_sheets": list(wb.sheetnames),
        "_restated_on_v2_agent": sorted({i[2] for i in issues if i[1] == "restated"}),
        "permissions": [perms[c] for c in order],
        "roles": [{"name": r, "codes": sorted(grants[r])} for r in roles_order],
        "issues": [i for i in issues if i[1] != "restated"],
    }


if __name__ == "__main__":
    model = extract(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_WB)
    with open(OUT, "w") as fh:
        json.dump({k: v for k, v in model.items() if k != "issues"}, fh, indent=1)
        fh.write("\n")
    print("wrote %s — %d permissions, %d roles" % (OUT, len(model["permissions"]), len(model["roles"])))
    for issue in model["issues"]:
        print("  issue:", issue)
